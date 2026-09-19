from PySide6.QtWidgets import (
    QDialog, QVBoxLayout, QHBoxLayout, QTableWidget, QTableWidgetItem, QHeaderView,
    QPushButton, QLineEdit, QLabel, QMessageBox, QDateEdit, QFileDialog
)
from PySide6.QtCore import QDate

from app import db
from app.services.plate_utils import format_plate, normalize_plate
from datetime import date, datetime


class SubscribersDialog(QDialog):
    """Abone Tanimlari: sabit ucretsiz/indirimli gecis hakki olan araclar."""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("Abone Tanimlari")
        self.resize(560, 420)

        layout = QVBoxLayout(self)

        form = QHBoxLayout()
        self.plate_input = QLineEdit()
        self.plate_input.setPlaceholderText("Plaka (34 GGG 01)")
        form.addWidget(self.plate_input)
        self.name_input = QLineEdit()
        self.name_input.setPlaceholderText("Ad Soyad / Firma")
        form.addWidget(self.name_input)
        layout.addLayout(form)

        date_row = QHBoxLayout()
        date_row.addWidget(QLabel("Gecerlilik Tarihi:"))
        self.valid_until = QDateEdit()
        self.valid_until.setCalendarPopup(True)
        self.valid_until.setDisplayFormat("yyyy-MM-dd")
        self.valid_until.setDate(QDate.currentDate().addMonths(1))
        date_row.addWidget(self.valid_until)
        self.unlimited_btn = QPushButton("Suresiz Yap")
        self.unlimited_btn.setCheckable(True)
        self.unlimited_btn.toggled.connect(self.valid_until.setDisabled)
        date_row.addWidget(self.unlimited_btn)
        date_row.addStretch()
        self.add_btn = QPushButton("Ekle / Guncelle")
        self.add_btn.setObjectName("primaryBtn")
        self.add_btn.clicked.connect(self.add_subscriber)
        date_row.addWidget(self.add_btn)
        layout.addLayout(date_row)

        self.table = QTableWidget(0, 4)
        self.table.setHorizontalHeaderLabels(["ID", "Plaka", "Isim Soyisim", "Gecerlilik"])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        layout.addWidget(self.table)

        btn_row = QHBoxLayout()
        template_btn = QPushButton("Excel Sablonunu Kaydet")
        template_btn.clicked.connect(self.save_template)
        btn_row.addWidget(template_btn)
        import_btn = QPushButton("Excel'den Abone Yukle")
        import_btn.setObjectName("primaryBtn")
        import_btn.clicked.connect(self.import_excel)
        btn_row.addWidget(import_btn)
        export_btn = QPushButton("Aboneleri Excel'e Indir")
        export_btn.clicked.connect(self.export_excel)
        btn_row.addWidget(export_btn)
        btn_row.addStretch()
        self.remove_btn = QPushButton("Secili Aboneyi Sil")
        self.remove_btn.clicked.connect(self.remove_selected)
        btn_row.addWidget(self.remove_btn)
        layout.addLayout(btn_row)

        self.reload()

    def save_template(self):
        target, _ = QFileDialog.getSaveFileName(self, "Abone Excel Sablonunu Kaydet",
                                                "abone_yukleme_sablonu.xlsx", "Excel (*.xlsx)")
        if not target:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            wb = Workbook()
            ws = wb.active
            ws.title = "Abone Yukleme"
            ws.append(["Plaka", "Isim Soyisim", "Gecerlilik Tarihi"])
            ws.append(["34 GGG 01", "Ornek Abone", "2027-12-31"])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2563EB")
            ws.column_dimensions["A"].width = 20
            ws.column_dimensions["B"].width = 32
            ws.column_dimensions["C"].width = 24
            wb.save(target)
            QMessageBox.information(self, "Kaydedildi", f"Sablon kaydedildi:\n{target}")
        except Exception as exc:
            QMessageBox.critical(self, "Excel Hatasi", str(exc))

    def import_excel(self):
        path, _ = QFileDialog.getOpenFileName(self, "Abone Excel Dosyasi Sec", "", "Excel (*.xlsx)")
        if not path:
            return
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            headers, header_row = [], None
            def header_key(value):
                return str(value or "").strip().lower().translate(str.maketrans("çğıöşü", "cgiosu"))
            for row_no, values in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
                candidate = [header_key(v) for v in values]
                if "plaka" in candidate and ("isim soyisim" in candidate or "ad/firma" in candidate):
                    headers, header_row = candidate, row_no
                    break
            if header_row is None:
                raise ValueError("Plaka, Isim Soyisim ve Gecerlilik Tarihi basliklari bulunamadi.")
            index = {name: i for i, name in enumerate(headers)}
            added, errors = 0, []
            name_header = "isim soyisim" if "isim soyisim" in index else "ad/firma"
            for row_no, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), header_row + 1):
                if not any(v not in (None, "") for v in row):
                    continue
                try:
                    raw_plate = row[index["plaka"]]
                    plate = normalize_plate(str(raw_plate or ""))
                    if not plate:
                        raise ValueError("gecersiz plaka")
                    valid = row[index["gecerlilik tarihi"]] if "gecerlilik tarihi" in index else None
                    if isinstance(valid, (datetime, date)):
                        valid = valid.strftime("%Y-%m-%d")
                    elif valid in (None, "", "Suresiz", "SÜRESİZ"):
                        valid = None
                    else:
                        valid = str(valid).strip()
                    def cell(name, default=""):
                        return row[index[name]] if name in index and index[name] < len(row) else default
                    name = str(cell(name_header) or "").strip()
                    if not name:
                        raise ValueError("isim soyisim bos")
                    db.add_subscriber(plate, name, "", valid)
                    added += 1
                except Exception as exc:
                    errors.append(f"Satir {row_no}: {exc}")
            wb.close()
            username = getattr(self.parent(), "current_user", {}).get("username", "sistem")
            db.audit(username, "ABONE_EXCEL_IMPORT", new_value=f"{added} kayit", note="; ".join(errors[:10]))
            self.reload()
            message = f"{added} abone eklendi/guncellendi."
            if errors:
                message += f"\n\n{len(errors)} satir atlandi:\n" + "\n".join(errors[:10])
            QMessageBox.information(self, "Excel Aktarimi", message)
        except Exception as exc:
            QMessageBox.critical(self, "Excel Hatasi", str(exc))

    def export_excel(self):
        target, _ = QFileDialog.getSaveFileName(self, "Aboneleri Excel'e Indir",
                                                "aboneler.xlsx", "Excel (*.xlsx)")
        if not target:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            wb = Workbook()
            ws = wb.active
            ws.title = "Aboneler"
            ws.append(["Plaka", "Isim Soyisim", "Gecerlilik Tarihi"])
            for sub in db.list_subscribers():
                ws.append([format_plate(sub["plate"]), sub.get("name") or "", sub.get("valid_until") or ""])
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="2563EB")
                cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions["A"].width = 20
            ws.column_dimensions["B"].width = 32
            ws.column_dimensions["C"].width = 24
            ws.freeze_panes = "A2"
            wb.save(target)
            username = getattr(self.parent(), "current_user", {}).get("username", "sistem")
            db.audit(username, "ABONE_EXCEL_EXPORT", new_value=f"{len(db.list_subscribers())} kayit")
            QMessageBox.information(self, "Excel Hazir", f"Aboneler indirildi:\n{target}")
        except Exception as exc:
            QMessageBox.critical(self, "Excel Hatasi", str(exc))

    def add_subscriber(self):
        plate = normalize_plate(self.plate_input.text())
        if not plate:
            QMessageBox.warning(self, "Hata", "Gecerli bir plaka giriniz.")
            return
        name = self.name_input.text().strip()
        if not name:
            QMessageBox.warning(self, "Hata", "Isim soyisim giriniz.")
            return
        valid_until = None if self.unlimited_btn.isChecked() else self.valid_until.date().toString("yyyy-MM-dd")
        db.add_subscriber(plate, name, "", valid_until)
        username = getattr(self.parent(), "current_user", {}).get("username", "sistem")
        db.audit(username, "ABONE_EKLENDI_GUNCELLENDI", target=plate,
                 new_value=f"{name} | {valid_until or 'Suresiz'}")
        self.plate_input.clear()
        self.name_input.clear()
        self.reload()

    def remove_selected(self):
        row = self.table.currentRow()
        if row < 0:
            return
        subscriber_id = int(self.table.item(row, 0).text())
        plate = self.table.item(row, 1).text()
        if QMessageBox.question(
                self, "Aboneyi Sil", f"{plate} plakali aboneyi silmek istiyor musunuz?") \
                != QMessageBox.StandardButton.Yes:
            return
        db.remove_subscriber(subscriber_id)
        username = getattr(self.parent(), "current_user", {}).get("username", "sistem")
        db.audit(username, "ABONE_SILINDI", target=normalize_plate(plate))
        self.reload()

    def reload(self):
        subs = db.list_subscribers()
        self.table.setRowCount(0)
        for s in subs:
            r = self.table.rowCount()
            self.table.insertRow(r)
            values = [str(s["id"]), format_plate(s["plate"]), s.get("name") or "-",
                      s.get("valid_until") or "Suresiz"]
            for c, v in enumerate(values):
                self.table.setItem(r, c, QTableWidgetItem(v))

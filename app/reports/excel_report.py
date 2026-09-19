"""Gunluk detayli Excel raporu (openpyxl)."""
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.config import REPORTS_DIR
from app.services.plate_utils import format_plate

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_sheet(wb, title, headers, rows):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for row in rows:
        ws.append(row)
    for col in ws.columns:
        width = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(40, max(12, width + 2))
    return ws


def build_excel(data: dict) -> str:
    wb = Workbook()
    wb.remove(wb.active)

    s = data["summary"]
    _write_sheet(wb, "Ozet", ["Alan", "Deger"], [
        ["Rapor Tarihi", data["report_date"]],
        ["Toplam Giris", s["toplam_giris"]],
        ["Toplam Cikis", s["toplam_cikis"]],
        ["Iceride Kalan", s["iceride_kalan"]],
        ["Normal Arac", s["normal_arac"]],
        ["Abone Arac", s["abone_arac"]],
        ["Toplam Tahsilat (TL)", s["toplam_tahsilat"]],
        ["Nakit (TL)", s["nakit"]],
        ["Kredi Karti (TL)", s["kredi_karti"]],
        ["Banka Karti (TL)", s["banka_karti"]],
        ["Havale/EFT (TL)", s["havale_eft"]],
        ["Iptal/Iade (TL)", s["iptal_iade"]],
        ["Manuel Bariyer Acma", s["manuel_bariyer"]],
        ["Okunamayan Plaka", s["okunamayan_plaka"]],
        ["Elle Duzeltilen Plaka", s["elle_duzeltilen"]],
    ])

    _write_sheet(wb, "Girisler", ["Plaka", "Giris", "Kapi", "Arac Tipi", "Abone", "Guven"],
                 [[format_plate(e["plaka"]), e["giris"], e["kapi"], e["arac_tipi"], e["abone"], e["guven"]]
                  for e in data["entries"]])
    _write_sheet(wb, "Cikislar",
                 ["Plaka", "Giris", "Cikis", "Sure", "Kapi", "Ucret (TL)", "Odeme", "Kasiyer"],
                 [[format_plate(e["plaka"]), e["giris"], e["cikis"], e["sure"], e["kapi"], e["ucret"],
                   e["odeme_yontemi"], e["kasiyer"]] for e in data["exits"]])
    _write_sheet(wb, "Iceride Kalanlar", ["Plaka", "Giris", "Sure", "Kapi", "Abone"],
                 [[format_plate(e["plaka"]), e["giris"], e["sure"], e["kapi"], e["abone"]] for e in data["inside"]])
    _write_sheet(wb, "Tahsilatlar", ["Islem No", "Plaka", "Tutar (TL)", "Yontem", "Kasiyer", "Saat", "Durum"],
                 [[p["islem_no"], format_plate(p["plaka"]), p["tutar"], p["yontem"], p["kasiyer"], p["saat"], p["durum"]]
                  for p in data["payments"]])
    _write_sheet(wb, "Manuel Islemler", ["Kullanici", "Kapi", "Plaka", "Neden", "Saat"],
                 [[m["kullanici"], m["kapi"], format_plate(m["plaka"]), m["neden"], m["saat"]]
                  for m in data["manual_ops"]])
    _write_sheet(wb, "Sistem Uyarilari", ["Uyari"], [[w] for w in data["warnings"]] or [["Uyari yok"]])

    path = os.path.join(REPORTS_DIR, f"Otopark_Gunluk_Rapor_{data['report_date']}.xlsx")
    wb.save(path)
    return path

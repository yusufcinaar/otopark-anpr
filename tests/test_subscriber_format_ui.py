"""Abone formu ve Excel aktarimi: izole veritabani, gercek Qt kontrolleri."""
import os
from pathlib import Path
import subprocess
import sys

import pytest


_PROBE = r'''
import os
import sys
from datetime import date, datetime
from pathlib import Path
from unittest.mock import patch

# Eski kurulum verisinin yeni test dizinine tasinmasini engelle.
data_dir = Path(os.environ['OTOPARK_DATA_DIR'])
data_dir.mkdir(parents=True, exist_ok=True)
(data_dir / 'otopark.db').touch()

from PySide6.QtCore import QDate
from PySide6.QtWidgets import QApplication
from app import db
from app.ui.subscribers_dialog import SubscribersDialog

case = sys.argv[1]
app = QApplication([])
db.init_db()
dialog = SubscribersDialog()
dialog.valid_until.setDate(QDate(2099, 12, 31))

def fill(plate='34 abc 123', name='Ali Veli'):
    dialog.plate_input.setText(plate)
    dialog.name_input.setText(name)

def table_values():
    return [[dialog.table.item(r, c).text() for c in range(4)]
            for r in range(dialog.table.rowCount())]

assert dialog.valid_until.displayFormat() == 'yyyy-MM-dd'
assert '34 GGG 01' in dialog.plate_input.placeholderText()

if case in ('dated', 'unlimited', 'back_to_date'):
    fill()
    if case != 'dated':
        dialog.unlimited_btn.click()
        assert dialog.unlimited_btn.isChecked()
        assert not dialog.valid_until.isEnabled()
    if case == 'back_to_date':
        dialog.unlimited_btn.click()
        assert not dialog.unlimited_btn.isChecked()
        assert dialog.valid_until.isEnabled()
        assert dialog.valid_until.date() == QDate(2099, 12, 31)
    dialog.add_btn.click()
    saved = db.list_subscribers()
    assert len(saved) == 1
    assert saved[0]['plate'] == '34ABC123'
    assert saved[0]['name'] == 'Ali Veli'
    expected_date = None if case == 'unlimited' else '2099-12-31'
    assert saved[0]['valid_until'] == expected_date
    assert table_values()[0][1:] == ['34 ABC 123', 'Ali Veli', expected_date or 'Suresiz']
    assert db.get_active_subscriber('34ABC123') is not None
    # Ayni plaka bosluk/tire farkiyla yeni kayit olusturmaz.
    fill('34-ABC-123', 'Guncel Isim')
    dialog.add_btn.click()
    assert len(db.list_subscribers()) == 1
    assert db.list_subscribers()[0]['name'] == 'Guncel Isim'
elif case == 'excel_roundtrip':
    from openpyxl import Workbook, load_workbook
    input_path = data_dir / 'input.xlsx'
    output_path = data_dir / 'export.xlsx'
    template_path = data_dir / 'template.xlsx'
    wb = Workbook()
    ws = wb.active
    headers = ['Plaka', 'Isim Soyisim', 'Gecerlilik Tarihi']
    ws.append(headers)
    ws.append(['34 abc 123', 'Ali Veli', '2099-12-31'])
    ws.append(['06-A-1234', 'Firma', None])
    ws.append(['35 AB 1234', 'Ayse', datetime(2098, 1, 2)])
    wb.save(input_path)
    with patch('app.ui.subscribers_dialog.QMessageBox.information'), \
         patch('app.ui.subscribers_dialog.QMessageBox.critical', side_effect=AssertionError), \
         patch('app.ui.subscribers_dialog.QFileDialog.getOpenFileName', return_value=(str(input_path), '')):
        dialog.import_excel()
        assert [r[1:] for r in table_values()] == [
            ['06 A 1234', 'Firma', 'Suresiz'],
            ['34 ABC 123', 'Ali Veli', '2099-12-31'],
            ['35 AB 1234', 'Ayse', '2098-01-02'],
        ]
        assert db.get_active_subscriber('06A1234')['valid_until'] is None
        with patch('app.ui.subscribers_dialog.QFileDialog.getSaveFileName', return_value=(str(output_path), '')):
            dialog.export_excel()
        exported = load_workbook(output_path, data_only=True)
        assert list(exported.active.values) == [
            tuple(headers), ('06 A 1234', 'Firma', None),
            ('34 ABC 123', 'Ali Veli', '2099-12-31'),
            ('35 AB 1234', 'Ayse', '2098-01-02'),
        ]
        exported.close()
        before_reimport = db.list_subscribers()
        with patch('app.ui.subscribers_dialog.QFileDialog.getOpenFileName', return_value=(str(output_path), '')):
            dialog.import_excel()
        assert db.list_subscribers() == before_reimport
        assert [s['plate'] for s in db.list_subscribers()] == ['06A1234', '34ABC123', '35AB1234']
        with patch('app.ui.subscribers_dialog.QFileDialog.getSaveFileName', return_value=(str(template_path), '')):
            dialog.save_template()
        template = load_workbook(template_path, data_only=True)
        assert list(template.active.values)[:2] == [tuple(headers), ('34 GGG 01', 'Ornek Abone', '2027-12-31')]
        template.close()
else:
    raise AssertionError(case)
dialog.close()
print('OK', case)
'''


@pytest.mark.parametrize("case", ["dated", "unlimited", "back_to_date", "excel_roundtrip"])
def test_subscriber_format_ui(case, tmp_path):
    result = subprocess.run(
        [sys.executable, "-c", _PROBE, case],
        cwd=Path(__file__).resolve().parents[1],
        env=dict(os.environ, QT_QPA_PLATFORM="offscreen",
                 OTOPARK_DATA_DIR=str(tmp_path / "subscriber-ui")),
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=20,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert f"OK {case}" in result.stdout
